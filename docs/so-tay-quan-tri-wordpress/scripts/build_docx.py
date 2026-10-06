#!/usr/bin/env python3
"""Build the A4 WordPress administration handbook from modular Markdown.

The renderer intentionally supports the small Markdown subset used by this
handbook.  It produces deterministic styles, a clickable static contents list,
bookmarks for every part/chapter/lesson, accessible image descriptions, and
page-number fields.
"""

from __future__ import annotations

import datetime
import io
import json
import re
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
OUTPUT = ROOT / "So-tay-quan-tri-website-WordPress.docx"
TOC_SOURCE = SRC / "00-muc-luc-de-xuat.md"
MANIFEST = ROOT / "screenshots" / "final" / "manifest.json"
THEME_STYLE = ROOT.parents[1] / "chinhtoa" / "style.css"

# Ảnh chụp ở độ nét 2x; manifest.json ghi bề rộng CSS (px) của từng ảnh. 1 px CSS in ra
# ~0,0178 cm (ảnh rộng 900 px = 16 cm, chữ giao diện 13 px cao ~2,3 mm — đọc được trên A4).
CM_PER_CSS_PX = 0.0178
MAX_IMAGE_W_CM = 16.0
MAX_IMAGE_H_CM = 21.5


def load_manifest() -> dict:
    try:
        return json.loads(MANIFEST.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


IMAGE_MANIFEST = load_manifest()


def theme_version() -> str:
    try:
        match = re.search(r"(?m)^Version:\s*(\S+)", THEME_STYLE.read_text(encoding="utf-8"))
        return match.group(1) if match else ""
    except OSError:
        return ""


def vietnamese_date(day: datetime.date) -> str:
    return f"{day.day:02d} tháng {day.month:02d} năm {day.year}"

SOURCE_ORDER = [
    "00-mo-dau.md",
    *[f"phan-01-chuong-{n:02d}.md" for n in range(1, 3)],
    *[f"phan-02-chuong-{n:02d}.md" for n in range(3, 7)],
    *[f"phan-03-chuong-{n:02d}.md" for n in range(7, 9)],
    *[f"phan-04-chuong-{n:02d}.md" for n in range(9, 11)],
    *[f"phan-05-chuong-{n:02d}.md" for n in range(11, 13)],
    *[f"phan-06-chuong-{n:02d}.md" for n in range(13, 17)],
    *[f"phan-07-chuong-{n:02d}.md" for n in range(17, 19)],
    *[f"phan-08-chuong-{n:02d}.md" for n in range(19, 21)],
    *[f"phan-09-chuong-{n:02d}.md" for n in range(21, 23)],
    *[f"phan-10-chuong-{n:02d}.md" for n in range(23, 25)],
    *[f"phu-luc-{letter}.md" for letter in "abcde"],
]

BLACK = "000000"
TEXT = "1F1F1F"
MUTED = "606060"
LINK = "135FA7"
BORDER = "D9D9D9"
TABLE_HEADER = "EAF0F6"
RED = "B91C1C"
GREEN = "16794A"
AMBER = "9A6700"


def set_cell_borders(cell, color: str = BORDER, width: str = "4") -> None:
    """Give every table cell a restrained, consistent grid."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), width)
        node.set(qn("w:color"), color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)


def set_cell_margins(
    cell,
    *,
    top: int = 90,
    start: int = 110,
    bottom: int = 90,
    end: int = 110,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_dxa_width(parent, tag: str, width: int) -> None:
    node = ensure_child(parent, tag)
    node.set(qn("w:type"), "dxa")
    node.set(qn("w:w"), str(width))


def apply_table_geometry(table, widths: list[int]) -> None:
    """Synchronize table, grid and cell widths for stable rendering."""
    total_width = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    tbl_pr = table._tbl.tblPr
    set_dxa_width(tbl_pr, "w:tblW", total_width)
    set_dxa_width(tbl_pr, "w:tblInd", 110)
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for col_index, width in enumerate(widths):
        table.columns[col_index].width = Twips(width)
    for row in table.rows:
        for col_index, cell in enumerate(row.cells):
            width = widths[col_index]
            cell.width = Twips(width)
            set_dxa_width(cell._tc.get_or_add_tcPr(), "w:tcW", width)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_heading(row) -> None:
    set_repeat_table_header(row)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def bookmark_name(kind: str, value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", value).strip("_").lower()
    return f"{kind}_{safe}"[:40]


def heading_bookmark(text: str) -> str | None:
    if text == "Mở đầu":
        return "part_00"
    if text == "Phần phụ lục":
        return "part_appendix"
    match = re.match(r"Phần\s+(\d+)\b", text)
    if match:
        return bookmark_name("part", match.group(1).zfill(2))
    match = re.match(r"Chương\s+(\d+)\b", text)
    if match:
        return bookmark_name("chapter", match.group(1).zfill(2))
    match = re.match(r"Phụ lục\s+([A-E])\b", text)
    if match:
        return bookmark_name("appendix", match.group(1))
    match = re.match(r"Bài\s+([A-Z0-9]+\.[0-9]+)\b", text)
    if match:
        return bookmark_name("lesson", match.group(1).replace(".", "_"))
    return None


class BookmarkCounter:
    def __init__(self) -> None:
        self.value = 1

    def add(self, paragraph, name: str | None) -> None:
        if not name:
            return
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.value))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.value))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)
        self.value += 1


def add_internal_link(
    paragraph,
    text: str,
    anchor: str,
    *,
    bold: bool = False,
    color: str = BLACK,
    underline: bool = False,
) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    if bold:
        b = OxmlElement("w:b")
        r_pr.append(b)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def add_inline(paragraph, text: str) -> None:
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.color.rgb = RGBColor.from_string("333333")
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def remove_style_paragraph_borders(style) -> None:
    """Remove template-era rules so hierarchy comes from type and spacing."""
    p_pr = style.element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    for name, size, before, after in [
        ("Title", 26, 0, 10),
        ("Subtitle", 11, 0, 6),
        ("Heading 1", 18, 16, 9),
        ("Heading 2", 15, 14, 7),
        ("Heading 3", 12, 12, 4),
        ("Heading 4", 11, 8, 3),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.font.underline = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        remove_style_paragraph_borders(style)

    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].paragraph_format.keep_with_next = True
    styles["Heading 3"].paragraph_format.keep_with_next = True
    styles["Heading 4"].paragraph_format.keep_with_next = True

    if "Caption Custom" not in styles:
        cap = styles.add_style("Caption Custom", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Caption Custom"]
    cap.font.name = "Arial"
    cap.font.size = Pt(9.5)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor.from_string(MUTED)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(7)
    cap.paragraph_format.keep_with_next = False

    if "Section Label" not in styles:
        label = styles.add_style("Section Label", WD_STYLE_TYPE.PARAGRAPH)
    else:
        label = styles["Section Label"]
    label.font.name = "Arial"
    label.font.size = Pt(11)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(BLACK)
    label.paragraph_format.space_before = Pt(6)
    label.paragraph_format.space_after = Pt(2)
    label.paragraph_format.keep_with_next = True

    if "Lesson Meta" not in styles:
        meta = styles.add_style("Lesson Meta", WD_STYLE_TYPE.PARAGRAPH)
    else:
        meta = styles["Lesson Meta"]
    meta.font.name = "Arial"
    meta.font.size = Pt(9.5)
    meta.font.bold = True
    meta.font.color.rgb = RGBColor.from_string(MUTED)
    meta.paragraph_format.space_after = Pt(4)
    meta.paragraph_format.keep_with_next = True

    if "Figure Reference" not in styles:
        figure_ref = styles.add_style("Figure Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        figure_ref = styles["Figure Reference"]
    figure_ref.font.name = "Arial"
    figure_ref.font.size = Pt(9.5)
    figure_ref.font.italic = True
    figure_ref.font.color.rgb = RGBColor.from_string(MUTED)
    figure_ref.paragraph_format.space_after = Pt(5)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Cm(0.62)
        style.paragraph_format.first_line_indent = Cm(-0.32)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.1

    custom_styles = [
        ("Cover Kicker", 9.5, True, False, BLACK),
        ("Cover Audience", 11, True, False, BLACK),
        ("Cover Meta", 9.5, False, False, MUTED),
        ("TOC Note", 9.5, False, True, MUTED),
        ("TOC Part", 11, True, False, BLACK),
        ("TOC Chapter", 11, False, False, BLACK),
        ("Table Header", 9.5, True, False, BLACK),
        ("Table Text", 9.5, False, False, TEXT),
    ]
    for name, size, bold, italic, color in custom_styles:
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.base_style = normal
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.space_after = Pt(3)

    styles["TOC Part"].paragraph_format.space_after = Pt(5)
    styles["TOC Chapter"].paragraph_format.left_indent = Cm(0.55)
    styles["Cover Kicker"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Cover Audience"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Cover Meta"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Table Header"].paragraph_format.space_after = Pt(0)
    styles["Table Text"].paragraph_format.space_after = Pt(0)


def configure_sections(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Sổ tay quản trị website WordPress")
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(BLACK)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("")
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    add_field(p, "PAGE")
    run = p.add_run(" / ")
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    add_field(p, "NUMPAGES")

    # The cover should remain completely quiet: no running header or folio.
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Cover Kicker")
    p.add_run("CHÍNH TÒA MEDIA")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.add_run("Sổ tay quản trị\nwebsite WordPress")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Hướng dẫn thao tác nhanh cho website sử dụng giao diện Chính Tòa Media")

    doc.add_paragraph()
    p = doc.add_paragraph(style="Cover Audience")
    p.add_run("Dành cho Biên tập viên · Quản lý · Hỗ trợ kỹ thuật")

    p = doc.add_paragraph(style="Cover Meta")
    version = theme_version()
    p.add_run("Kiểm chứng trên WordPress 7.1 và Chính Tòa Media" + (f" {version}" if version else ""))

    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Cover Meta")
    p.add_run("Phiên bản ngày " + vietnamese_date(datetime.date.today()))

    doc.add_page_break()


def parse_toc_entries() -> list[tuple[int, str, str, str]]:
    entries: list[tuple[int, str, str, str]] = []
    for line in TOC_SOURCE.read_text(encoding="utf-8").splitlines():
        if line.startswith("# Phần "):
            text = line[2:].strip()
            anchor = heading_bookmark(text) or ""
            entries.append((0, text, anchor, "part"))
        elif line.startswith("## Chương ") or line.startswith("## Phụ lục "):
            text = line[3:].strip()
            anchor = heading_bookmark(text) or ""
            entries.append((1, text, anchor, "chapter"))
        # The quick contents deliberately stops at chapter level. Lesson-level
        # bookmarks remain available in Word's navigation pane without turning
        # the opening pages into a dense 177-line index.
    return entries


def add_static_toc(doc: Document) -> None:
    title = doc.add_paragraph(style="Heading 1")
    title.add_run("Mục lục nhanh")
    note = doc.add_paragraph(style="TOC Note")
    note.add_run("Chọn một phần hoặc chương để chuyển nhanh đến nội dung cần xem.")

    for level, text, anchor, kind in parse_toc_entries():
        p = doc.add_paragraph(style="TOC Part" if level == 0 else "TOC Chapter")
        p.paragraph_format.keep_with_next = False
        add_internal_link(
            p,
            text,
            anchor,
            bold=(level == 0),
            color=BLACK,
            underline=False,
        )

    doc.add_page_break()


def add_heading(doc: Document, level: int, text: str, bookmarks: BookmarkCounter) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    add_inline(p, text)
    bookmarks.add(p, heading_bookmark(text))


def add_image(doc: Document, source_file: Path, alt: str) -> None:
    """Chèn ảnh theo kích thước thật: không phóng ảnh nhỏ, không ép mọi ảnh rộng 16 cm.

    Ảnh được nhúng dạng JPEG chất lượng cao (không giảm mẫu màu) để tệp DOCX gọn mà
    viền đỏ, chữ nhỏ vẫn sắc. Cùng một ảnh dùng ở nhiều bài chỉ lưu một lần.
    """
    with Image.open(source_file) as image:
        px_w, px_h = image.size
        css_w = IMAGE_MANIFEST.get(source_file.stem, {}).get("w") or px_w / 2
        width_cm = min(MAX_IMAGE_W_CM, css_w * CM_PER_CSS_PX)
        height_cm = width_cm * px_h / px_w
        if height_cm > MAX_IMAGE_H_CM:
            height_cm = MAX_IMAGE_H_CM
            width_cm = height_cm * px_w / px_h
        buffer = io.BytesIO()
        image.convert("RGB").save(buffer, "JPEG", quality=90, subsampling=0, optimize=True)
        buffer.seek(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(buffer, width=Cm(width_cm), height=Cm(height_cm))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt[:100])


def add_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Give descriptive columns more room while preserving a useful minimum.
    lengths = []
    for col_index in range(cols):
        column_values = [row[col_index] if col_index < len(row) else "" for row in rows]
        plain_lengths = [len(re.sub(r"[*`]", "", value)) for value in column_values]
        lengths.append(max(8, min(max(plain_lengths, default=8), 60)))
    min_width = 3.0
    total_width = 16.7
    remaining = total_width - (min_width * cols)
    weight_total = sum(lengths)
    width_cm = [min_width + remaining * (weight / weight_total) for weight in lengths]
    total_width_dxa = int(round(Cm(total_width).twips))
    widths = [int(round(Cm(value).twips)) for value in width_cm]
    widths[-1] += total_width_dxa - sum(widths)

    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = "Table Header" if row_index == 0 else "Table Text"
            if col_index < len(row):
                add_inline(p, row[col_index].strip())
            set_cell_margins(cell)
            set_cell_borders(cell)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
    apply_table_geometry(table, widths)
    set_repeat_heading(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def parse_source(doc: Document, path: Path, bookmarks: BookmarkCounter, seen_parts: set[str]) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    just_added_part = False
    current_section: str | None = None
    section_content_count = 0
    current_risk = ""
    error_bullets_shown = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            if title == "Sổ tay quản trị website WordPress":
                index += 1
                continue
            if title not in seen_parts:
                if len(doc.paragraphs) > 0:
                    doc.add_page_break()
                add_heading(doc, 1, title, bookmarks)
                seen_parts.add(title)
                just_added_part = True
            current_section = None
            index += 1
            continue

        if line.startswith("## "):
            title = line[3:].strip()
            add_heading(doc, 2, title, bookmarks)
            just_added_part = False
            current_section = None
            index += 1
            continue

        if line.startswith("### "):
            add_heading(doc, 3, line[4:].strip(), bookmarks)
            just_added_part = False
            current_section = None
            current_risk = ""
            error_bullets_shown = 0
            index += 1
            continue

        if line.startswith("#### "):
            current_section = line[5:].strip()
            section_content_count = 0
            if current_section == "Các bước thực hiện":
                p = doc.add_paragraph(style="Section Label")
                p.add_run("Thực hiện")
            elif current_section == "Nếu gặp lỗi":
                p = doc.add_paragraph(style="Section Label")
                p.add_run("Nếu chưa đúng")
            index += 1
            continue

        image_match = re.fullmatch(r"!\[(.*)\]\((.*)\)", line)
        if image_match:
            alt, relative = image_match.groups()
            source_file = (path.parent / relative).resolve()
            if not source_file.is_file():
                raise FileNotFoundError(source_file)
            add_image(doc, source_file, alt)
            section_content_count += 1
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1].strip()):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                current = lines[index].strip()
                if not is_table_separator(current):
                    table_rows.append([cell.strip() for cell in current.strip("|").split("|")])
                index += 1
            add_table(doc, table_rows)
            continue

        if re.match(r"^-\s+", line):
            if current_section == "Nếu gặp lỗi":
                # Bản mới: mỗi bài có 2–4 tình huống riêng, hiện đủ cho người đọc.
                if error_bullets_shown >= 4:
                    index += 1
                    continue
                error_bullets_shown += 1
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, re.sub(r"^-\s+", "", line))
            else:
                bullet_leads = {
                    "Mục đích": "Mục tiêu",
                    "Trước khi bắt đầu": "Chuẩn bị",
                    "Kết quả": "Kết quả",
                    "Lưu ý": "Cần nhớ",
                    "Không nên làm": "Tránh",
                }
                if current_section in bullet_leads and section_content_count == 0:
                    # Mục viết bằng gạch đầu dòng vẫn cần nhãn đầu mục như mục viết thành đoạn.
                    label = doc.add_paragraph()
                    label.paragraph_format.space_after = Pt(1)
                    label.paragraph_format.keep_with_next = True
                    lead = label.add_run(f"{bullet_leads[current_section]}:")
                    lead.bold = True
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, re.sub(r"^-\s+", "", line))
            section_content_count += 1
            index += 1
            continue

        number_match = re.match(r"^(\d+)\.\s+(.*)", line)
        if number_match:
            # Keep the source number. Word's shared List Number definition
            # otherwise continues numbering across unrelated lessons.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.62)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            step_text = re.sub(
                r"^\*\*Bước\s+\d+:\*\*\s*",
                "",
                number_match.group(2),
            )
            number = p.add_run(f"{number_match.group(1)}.  ")
            number.bold = True
            add_inline(p, step_text)
            section_content_count += 1
            index += 1
            continue

        if line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            p = doc.add_paragraph(style="Caption Custom")
            add_inline(p, line)
            section_content_count += 1
            index += 1
            continue

        if line.startswith("**Nhãn:**"):
            label_match = re.match(r"\*\*Nhãn:\*\*\s+(.+?)\s+·\s+(.+)$", line)
            p = doc.add_paragraph(style="Lesson Meta")
            if label_match:
                audience, safety = label_match.groups()
                audience_short = {
                    "Bắt buộc học": "BẮT BUỘC",
                    "Nên biết": "NÊN BIẾT",
                    "Kỹ thuật dành cho quản trị viên": "QUẢN LÝ / KỸ THUẬT",
                }.get(audience, audience.upper())
                risk_color = MUTED
                safety_short = safety
                if "🟢" in safety:
                    current_risk = "green"
                    risk_color = GREEN
                    safety_short = "AN TOÀN"
                elif "🟡" in safety:
                    current_risk = "amber"
                    risk_color = AMBER
                    safety_short = "CẨN THẬN"
                elif "🔴" in safety:
                    current_risk = "red"
                    risk_color = RED
                    safety_short = "KHÔNG TỰ THAY ĐỔI"
                audience_run = p.add_run(audience_short)
                audience_run.font.color.rgb = RGBColor.from_string(MUTED)
                separator = p.add_run("   ·   ")
                separator.font.color.rgb = RGBColor.from_string(MUTED)
                risk_dot = p.add_run("● ")
                risk_dot.font.color.rgb = RGBColor.from_string(risk_color)
                safety_run = p.add_run(safety_short)
                safety_run.font.color.rgb = RGBColor.from_string(MUTED)
            else:
                add_inline(p, line)
            index += 1
            continue


        if current_section == "Ảnh minh họa":
            if line.lower().startswith("không cần ảnh"):
                index += 1
                continue
            p = doc.add_paragraph(style="Figure Reference")
            add_inline(p, line)
            section_content_count += 1
            index += 1
            continue

        lead_labels = {
            "Mục đích": "Mục tiêu",
            "Trước khi bắt đầu": "Chuẩn bị",
            "Kết quả": "Kết quả",
            "Lưu ý": "Cần nhớ",
            "Không nên làm": "Tránh",
        }
        if current_section in lead_labels:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2.5)
            if current_section in {
                "Mục đích",
                "Trước khi bắt đầu",
                "Lưu ý",
                "Không nên làm",
            }:
                p.paragraph_format.keep_with_next = True
            if section_content_count == 0:
                lead = p.add_run(f"{lead_labels[current_section]} — ")
                lead.bold = True
            add_inline(p, line)
            section_content_count += 1
            index += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, line)
        section_content_count += 1
        index += 1


def set_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Sổ tay quản trị website WordPress"
    props.subject = "Hướng dẫn quản trị website sử dụng giao diện Chính Tòa Media"
    props.author = "Chính Tòa Media"
    props.keywords = "WordPress, Gutenberg, Chính Tòa Media, quản trị website"
    props.comments = "Biên soạn từ giao diện thực tế (WordPress 7.1) và mã nguồn theme Chính Tòa Media."


def enable_field_updates(doc: Document) -> None:
    """Ask Word/LibreOffice to refresh PAGE and NUMPAGES fields on open."""
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> int:
    missing = [name for name in SOURCE_ORDER if not (SRC / name).is_file()]
    if missing:
        print("Missing source files:", *missing, sep="\n- ", file=sys.stderr)
        return 1

    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    set_document_properties(doc)
    bookmarks = BookmarkCounter()

    add_cover(doc)
    add_static_toc(doc)
    add_heading(doc, 1, "Mở đầu", bookmarks)

    seen_parts: set[str] = set()
    for name in SOURCE_ORDER:
        parse_source(doc, SRC / name, bookmarks, seen_parts)

    enable_field_updates(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
